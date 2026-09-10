"""Genera data/comuni.json dal permalink ufficiale ISTAT."""

from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl


REGION_ALIASES = {
    "Trentino-Alto Adige/Südtirol": "Trentino-Alto Adige",
    "Valle d'Aosta/Vallée d'Aoste": "Valle d'Aosta",
}
EXPECTED_REGIONS = {
    "Abruzzo",
    "Basilicata",
    "Calabria",
    "Campania",
    "Emilia-Romagna",
    "Friuli-Venezia Giulia",
    "Lazio",
    "Liguria",
    "Lombardia",
    "Marche",
    "Molise",
    "Piemonte",
    "Puglia",
    "Sardegna",
    "Sicilia",
    "Toscana",
    "Trentino-Alto Adige",
    "Umbria",
    "Valle d'Aosta",
    "Veneto",
}
EXPECTED_TOTAL = 7_894


def main() -> None:
    if len(sys.argv) != 2:
        raise RuntimeError("Percorso XLSX ISTAT mancante")
    workbook = openpyxl.load_workbook(Path(sys.argv[1]), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header = [str(value or "").strip() for value in next(sheet.iter_rows(values_only=True))]
    name_idx = header.index("Denominazione in italiano")
    region_idx = header.index("Denominazione Regione")

    grouped: dict[str, set[str]] = defaultdict(set)
    for row in sheet.iter_rows(values_only=True):
        name = str(row[name_idx] or "").strip()
        raw_region = str(row[region_idx] or "").strip()
        region = REGION_ALIASES.get(raw_region, raw_region)
        if name and region in EXPECTED_REGIONS:
            grouped[region].add(name)

    if set(grouped) != EXPECTED_REGIONS:
        missing = sorted(EXPECTED_REGIONS - set(grouped))
        raise RuntimeError(f"Regioni ISTAT mancanti: {missing}")
    total = sum(len(names) for names in grouped.values())
    if total != EXPECTED_TOTAL:
        raise RuntimeError(f"Totale comuni inatteso: {total}, attesi {EXPECTED_TOTAL}")

    output = {
        region: sorted(grouped[region], key=str.casefold)
        for region in sorted(EXPECTED_REGIONS, key=str.casefold)
    }
    target = Path.cwd() / "data" / "comuni.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(f".tmp.{os.getpid()}.json")
    temporary.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(target)

    print(f"ISTAT: {total} comuni in {len(output)} regioni")
    for region, names in output.items():
        print(f"  {region}: {len(names)}")
    print(f"Salvato: {target}")


if __name__ == "__main__":
    main()
